"""File parsing utilities for PDF and DOCX documents (PRD Section 1.3.2 & 2.2)."""

import io
from typing import Literal

import docx
import pdfplumber
from fastapi import HTTPException, UploadFile

MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB

ALLOWED_MIME_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/msword",
}


def parse_pdf(file_bytes: bytes) -> str:
    """Extract text and flattened tables from a PDF using pdfplumber."""
    extracted_parts: list[str] = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                extracted_parts.append(text.strip())

            tables = page.extract_tables()
            for table in tables:
                table_lines: list[str] = []
                for row in table:
                    # Clean None/empty cell values and join with pipe separator
                    row_str = " | ".join(
                        (cell.strip() if cell else "") for cell in row
                    )
                    if row_str.strip(" |"):
                        table_lines.append(row_str)
                if table_lines:
                    extracted_parts.append("\n".join(table_lines))

    return "\n\n".join(extracted_parts)


def parse_docx(file_bytes: bytes) -> str:
    """Extract paragraphs and flattened tables from a DOCX file using python-docx."""
    doc = docx.Document(io.BytesIO(file_bytes))
    extracted_parts: list[str] = []

    for paragraph in doc.paragraphs:
        if paragraph.text and paragraph.text.strip():
            extracted_parts.append(paragraph.text.strip())

    for table in doc.tables:
        table_lines: list[str] = []
        for row in table.rows:
            row_str = " | ".join(cell.text.strip() for cell in row.cells)
            if row_str.strip(" |"):
                table_lines.append(row_str)
        if table_lines:
            extracted_parts.append("\n".join(table_lines))

    return "\n\n".join(extracted_parts)


def validate_and_parse_file(
    file: UploadFile, content: bytes
) -> tuple[str, Literal["pdf", "docx"]]:
    """Validate file size, extension, MIME type, and return parsed text and source type."""
    if len(content) > MAX_FILE_SIZE_BYTES:
        raise HTTPException(
            status_code=400,
            detail="File size exceeds the 10MB limit.",
        )

    filename = (file.filename or "").lower()
    content_type = (file.content_type or "").lower()

    if filename.endswith(".pdf") or content_type == "application/pdf":
        try:
            parsed_text = parse_pdf(content)
            return parsed_text, "pdf"
        except Exception as err:
            raise HTTPException(
                status_code=400,
                detail=f"Failed to parse PDF file: {err}",
            ) from err

    elif (
        filename.endswith(".docx")
        or content_type in ALLOWED_MIME_TYPES
    ):
        try:
            parsed_text = parse_docx(content)
            return parsed_text, "docx"
        except Exception as err:
            raise HTTPException(
                status_code=400,
                detail=f"Failed to parse DOCX file: {err}",
            ) from err
    else:
        raise HTTPException(
            status_code=400,
            detail="Unsupported file type. Only PDF (.pdf) and Word (.docx) files are supported.",
        )
